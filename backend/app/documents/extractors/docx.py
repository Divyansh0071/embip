"""
DOCX Document Extractor using python-docx.
Extracts paragraphs and tables text blocks from Microsoft Word documents.
"""

import io
import docx
from app.documents.extractors.base import DocumentExtractor, ExtractedDocument, TextBlock


class DOCXExtractor(DocumentExtractor):
    """Extractor for Microsoft Word (.docx) files."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        doc = docx.Document(io.BytesIO(file_bytes))
        blocks = []

        # 1. Paragraphs
        paragraph_texts = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraph_texts.append(text)

        if paragraph_texts:
            blocks.append(
                TextBlock(
                    text="\n\n".join(paragraph_texts),
                    metadata={"source_filename": filename, "section": "paragraphs"},
                )
            )

        # 2. Tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cell_texts:
                    table_rows.append(" | ".join(cell_texts))

            if table_rows:
                blocks.append(
                    TextBlock(
                        text=f"[Table {t_idx + 1}]\n" + "\n".join(table_rows),
                        row_count=len(table_rows),
                        metadata={"source_filename": filename, "table_index": t_idx + 1},
                    )
                )

        return ExtractedDocument(
            filename=filename,
            file_type="docx",
            text_blocks=blocks,
            extracted_metadata={"paragraph_count": len(paragraph_texts), "table_count": len(doc.tables)},
        )

"""
XLSX Excel Document Extractor using openpyxl.
Extracts spreadsheet workbooks and worksheets into structured text representations.
"""

import io
import openpyxl
from app.documents.extractors.base import DocumentExtractor, ExtractedDocument, TextBlock


class XLSXExtractor(DocumentExtractor):
    """Extractor for Microsoft Excel (.xlsx) spreadsheet workbooks."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        blocks = []
        total_sheets = len(workbook.sheetnames)

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet_rows = []

            for row in sheet.iter_rows(values_only=True):
                cell_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if cell_vals:
                    sheet_rows.append(" | ".join(cell_vals))

            if sheet_rows:
                blocks.append(
                    TextBlock(
                        text=f"[Sheet: {sheet_name}]\n" + "\n".join(sheet_rows),
                        sheet_name=sheet_name,
                        row_count=len(sheet_rows),
                        metadata={
                            "source_filename": filename,
                            "sheet_name": sheet_name,
                            "row_count": len(sheet_rows),
                        },
                    )
                )

        return ExtractedDocument(
            filename=filename,
            file_type="xlsx",
            text_blocks=blocks,
            total_sheets=total_sheets,
            extracted_metadata={"total_sheets": total_sheets, "sheet_names": workbook.sheetnames},
        )
